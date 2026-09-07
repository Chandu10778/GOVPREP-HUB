from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor

# Create presentation
prs = Presentation()
prs.slide_width = Inches(10)
prs.slide_height = Inches(7.5)

# Define color scheme
PRIMARY_COLOR = RGBColor(59, 130, 246)  # Blue
SECONDARY_COLOR = RGBColor(139, 92, 246)  # Purple
ACCENT_COLOR = RGBColor(16, 185, 129)  # Green
DARK_TEXT = RGBColor(31, 41, 55)
LIGHT_TEXT = RGBColor(255, 255, 255)

def add_title_slide(prs, title, subtitle):
    """Add title slide"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # Blank layout
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = PRIMARY_COLOR
    
    # Add title
    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(2.5), Inches(9), Inches(1.5))
    title_frame = title_box.text_frame
    title_frame.word_wrap = True
    p = title_frame.paragraphs[0]
    p.text = title
    p.font.size = Pt(60)
    p.font.bold = True
    p.font.color.rgb = LIGHT_TEXT
    p.alignment = PP_ALIGN.CENTER
    
    # Add subtitle
    subtitle_box = slide.shapes.add_textbox(Inches(0.5), Inches(4.2), Inches(9), Inches(1.5))
    subtitle_frame = subtitle_box.text_frame
    subtitle_frame.word_wrap = True
    p = subtitle_frame.paragraphs[0]
    p.text = subtitle
    p.font.size = Pt(28)
    p.font.color.rgb = LIGHT_TEXT
    p.alignment = PP_ALIGN.CENTER

def add_content_slide(prs, title, content_points, bg_color=RGBColor(255, 255, 255)):
    """Add content slide with title and bullet points"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # Blank layout
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = bg_color
    
    # Add title bar
    title_shape = slide.shapes.add_shape(1, Inches(0), Inches(0), Inches(10), Inches(1))
    title_shape.fill.solid()
    title_shape.fill.fore_color.rgb = PRIMARY_COLOR
    title_shape.line.color.rgb = PRIMARY_COLOR
    
    # Add title text
    title_frame = title_shape.text_frame
    p = title_frame.paragraphs[0]
    p.text = title
    p.font.size = Pt(44)
    p.font.bold = True
    p.font.color.rgb = LIGHT_TEXT
    p.alignment = PP_ALIGN.LEFT
    title_frame.margin_left = Inches(0.5)
    title_frame.margin_top = Inches(0.15)
    
    # Add content
    content_box = slide.shapes.add_textbox(Inches(0.8), Inches(1.5), Inches(8.4), Inches(5.5))
    text_frame = content_box.text_frame
    text_frame.word_wrap = True
    
    for i, point in enumerate(content_points):
        if i == 0:
            p = text_frame.paragraphs[0]
        else:
            p = text_frame.add_paragraph()
        p.text = point
        p.font.size = Pt(22)
        p.font.color.rgb = DARK_TEXT
        p.space_before = Pt(12)
        p.space_after = Pt(12)
        p.level = 0

def add_two_column_slide(prs, title, left_content, right_content):
    """Add slide with two columns"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = RGBColor(255, 255, 255)
    
    # Add title bar
    title_shape = slide.shapes.add_shape(1, Inches(0), Inches(0), Inches(10), Inches(0.9))
    title_shape.fill.solid()
    title_shape.fill.fore_color.rgb = SECONDARY_COLOR
    title_shape.line.color.rgb = SECONDARY_COLOR
    
    title_frame = title_shape.text_frame
    p = title_frame.paragraphs[0]
    p.text = title
    p.font.size = Pt(40)
    p.font.bold = True
    p.font.color.rgb = LIGHT_TEXT
    p.alignment = PP_ALIGN.LEFT
    title_frame.margin_left = Inches(0.5)
    title_frame.margin_top = Inches(0.15)
    
    # Left column
    left_box = slide.shapes.add_textbox(Inches(0.5), Inches(1.2), Inches(4.3), Inches(6))
    left_frame = left_box.text_frame
    left_frame.word_wrap = True
    for i, point in enumerate(left_content):
        if i == 0:
            p = left_frame.paragraphs[0]
        else:
            p = left_frame.add_paragraph()
        p.text = "• " + point
        p.font.size = Pt(18)
        p.font.color.rgb = DARK_TEXT
        p.space_after = Pt(10)
    
    # Right column
    right_box = slide.shapes.add_textbox(Inches(5.2), Inches(1.2), Inches(4.3), Inches(6))
    right_frame = right_box.text_frame
    right_frame.word_wrap = True
    for i, point in enumerate(right_content):
        if i == 0:
            p = right_frame.paragraphs[0]
        else:
            p = right_frame.add_paragraph()
        p.text = "• " + point
        p.font.size = Pt(18)
        p.font.color.rgb = DARK_TEXT
        p.space_after = Pt(10)

# Slide 1: Title Slide
add_title_slide(prs, "AI-Powered Learning Platform", "Revolutionizing Education Through Intelligent Technology")

# Slide 2: Problem Statement
add_content_slide(prs, "The Challenge", [
    "🎯 Traditional learning lacks personalization and engagement",
    "📚 Students struggle with diverse learning materials and topics",
    "🤖 Limited access to intelligent tutoring and real-time assistance",
    "📊 Difficulty tracking progress and maintaining motivation"
])

# Slide 3: Our Solution
add_content_slide(prs, "Our Solution", [
    "✨ Intelligent AI-powered chatbot providing personalized learning support",
    "📖 Comprehensive study materials with interactive content",
    "✅ Task management and assignment tracking system",
    "🔥 Streak system to motivate consistent learning",
    "🌐 Modern, user-friendly web application interface"
], RGBColor(240, 249, 255))

# Slide 4: Key Features
add_two_column_slide(prs, "Key Features", [
    "AI Chatbot Assistant",
    "Smart Study Materials",
    "Interactive Dashboard",
    "User Profiles & Progress",
], [
    "Task Management",
    "Previous Year Questions",
    "Current Affairs Section",
    "Performance Streaks"
])

# Slide 5: Platform Architecture
add_content_slide(prs, "Platform Architecture", [
    "🖥️ Frontend: React.js with Tailwind CSS for responsive UI",
    "⚙️ Backend: Flask Python with REST API architecture",
    "🗄️ Database: SQLAlchemy ORM for data management",
    "🧠 AI Integration: Groq API for advanced language capabilities"
])

# Slide 6: User Experience Flow
add_two_column_slide(prs, "User Experience", [
    "User Authentication",
    "Dashboard Overview",
    "Browse Study Materials",
    "Access Chatbot",
], [
    "Manage Tasks",
    "Track Progress",
    "View Streaks",
    "Update Profile"
])

# Slide 7: Technology Stack
add_two_column_slide(prs, "Technology Stack", [
    "Frontend Technologies:",
    "• React.js",
    "• React Router",
    "• Tailwind CSS",
    "• Lucide Icons",
], [
    "Backend Technologies:",
    "• Flask Framework",
    "• SQLAlchemy ORM",
    "• Flask-CORS",
    "• Groq AI API"
])

# Slide 8: Core Features Deep Dive
add_content_slide(prs, "AI Chatbot Capabilities", [
    "💬 Real-time conversational learning assistance",
    "🎓 Subject-specific tutoring and explanations",
    "📝 Problem-solving guidance and example solutions",
    "🔄 Multi-turn conversation support",
    "⚡ Powered by Groq's fast API"
])

# Slide 9: Study Materials System
add_content_slide(prs, "Study Materials System", [
    "📚 Curated learning resources and content",
    "🏷️ Organized by subjects and topics",
    "📈 Comprehensive coverage of exam materials",
    "🔗 Integrated with chatbot for deeper learning",
    "✏️ Easy navigation and search functionality"
], RGBColor(240, 249, 255))

# Slide 10: Task & Progress Management
add_content_slide(prs, "Task Management & Progress", [
    "✅ Create and track assignments and tasks",
    "📅 Deadline management and reminders",
    "🎯 Subject-based task organization",
    "🏆 Streak counter for motivation",
    "📊 Visual progress tracking dashboard"
])

# Slide 11: Impact & Benefits
add_two_column_slide(prs, "Impact & Benefits", [
    "For Students:",
    "• Personalized learning experience",
    "• 24/7 AI tutoring access",
    "• Better organization",
    "• Increased engagement",
], [
    "For Institutions:",
    "• Student retention improvement",
    "• Better learning outcomes",
    "• Modern tech infrastructure",
    "• Data-driven insights"
])

# Slide 12: Roadmap & Future Enhancements
add_content_slide(prs, "Future Roadmap", [
    "🚀 Mobile app development (iOS/Android)",
    "📊 Advanced analytics and reporting",
    "🌍 Multi-language support",
    "🤝 Peer-to-peer learning features",
    "🎮 Gamification elements and achievements"
], RGBColor(240, 249, 255))

# Slide 13: Call to Action
slide = prs.slides.add_slide(prs.slide_layouts[6])
background = slide.background
fill = background.fill
fill.solid()
fill.fore_color.rgb = SECONDARY_COLOR

# Add main text
main_box = slide.shapes.add_textbox(Inches(0.5), Inches(2.5), Inches(9), Inches(2))
main_frame = main_box.text_frame
main_frame.word_wrap = True
p = main_frame.paragraphs[0]
p.text = "Transform Learning Today"
p.font.size = Pt(54)
p.font.bold = True
p.font.color.rgb = LIGHT_TEXT
p.alignment = PP_ALIGN.CENTER

# Add subtitle
sub_box = slide.shapes.add_textbox(Inches(0.5), Inches(4.2), Inches(9), Inches(2))
sub_frame = sub_box.text_frame
sub_frame.word_wrap = True
p = sub_frame.paragraphs[0]
p.text = "Join us in revolutionizing education"
p.font.size = Pt(32)
p.font.color.rgb = LIGHT_TEXT
p.alignment = PP_ALIGN.CENTER

# Save presentation
output_path = r"c:\Users\chand\OneDrive\Desktop\project\AI_Learning_Platform_Presentation.pptx"
prs.save(output_path)
print(f"[SUCCESS] Presentation created successfully: {output_path}")
